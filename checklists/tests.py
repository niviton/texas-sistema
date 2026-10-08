import base64
import io
import os
import shutil
import tempfile

import docx
from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from PIL import Image

from accounts.models import User
from veiculos.models import Supervisor

from .importador import ler_formulario
from .models import Ativo, Execucao, Item, Modelo, Revisao, TipoAtivo
from .pdf import render_execucao_pdf

TMP_MEDIA = tempfile.mkdtemp()


def formulario_docx(codigo='TCB-OTB-99', revisao='01'):
    """Monta um .docx no formato dos formulários TCB-OTB (cabeçalho bilíngue + tabelas Item | OK | NOK)."""
    d = docx.Document()
    h = d.sections[0].header
    for texto in ['Título', 'Title', 'CHECKLIST DE PRÉ-USO – GERADOR DIESEL', 'Documento', 'Document', 'Revisão', 'Review',
                  'Página', 'Pages', codigo, revisao, '1 de 1', 'Preparado por', 'Prepared by', 'Revisado por', 'Review by',
                  'Data', 'Date', 'Ana Exemplo', 'Bruno Exemplo', '20/05/2025']:
        h.add_paragraph(texto)
    d.add_paragraph('CHECKLIST DE PRÉ-USO')
    d.add_paragraph('INFORMAÇÕES DO EQUIPAMENTO')
    info = d.add_table(rows=2, cols=2)
    info.cell(0, 0).text, info.cell(1, 0).text = 'IDENTIFICAÇÃO / N. SÉRIE:', 'MODELO:'
    for titulo, itens in [('VERIFICAÇÃO DOS EPIs', ['Óculos de segurança?', 'Luvas?']),
                          ('VERIFICAÇÃO DO EQUIPAMENTO', ['Nível de óleo ok?', 'Cabos sem danos?', 'Painel sem alarmes?'])]:
        d.add_paragraph(titulo)
        t = d.add_table(rows=1 + len(itens), cols=3)
        t.cell(0, 0).text, t.cell(0, 1).text, t.cell(0, 2).text = 'Item', 'OK', 'NOK'
        for i, texto in enumerate(itens, start=1):
            t.cell(i, 0).text = texto
    d.add_paragraph('INFORMAÇÕES ADICIONAIS')
    obs = d.add_table(rows=1, cols=1)
    obs.cell(0, 0).text = 'OBSERVAÇÕES'
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _jpeg(nome='foto.jpg'):
    buf = io.BytesIO()
    Image.new('RGB', (60, 40), 'gray').save(buf, 'JPEG')
    return SimpleUploadedFile(nome, buf.getvalue(), content_type='image/jpeg')


def _assinatura():
    buf = io.BytesIO()
    img = Image.new('RGBA', (200, 80), (0, 0, 0, 0))
    for x in range(10, 190):
        img.putpixel((x, 40), (13, 27, 46, 255))
    img.save(buf, 'PNG')
    return 'data:image/png;base64,' + base64.b64encode(buf.getvalue()).decode()


@override_settings(MEDIA_ROOT=TMP_MEDIA, EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend', VEICULOS_EMAIL_ASYNC=False)
class ChecklistsTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(TMP_MEDIA, ignore_errors=True)

    def setUp(self):
        self.admin = User.objects.create_superuser(email='admin@x.com', password='x', full_name='Admin Geral')
        self.tec = User.objects.create_user(email='tec@x.com', password='x', full_name='Tiago Técnico', is_approved=True, role=User.ROLE_TECNICO)
        self.outro = User.objects.create_user(email='outro@x.com', password='x', full_name='Outro Técnico', is_approved=True, role=User.ROLE_TECNICO)
        Supervisor.objects.create(name='Sup', email='sup@x.com')
        self.client.force_login(self.admin)
        r = self.client.post(reverse('checklists:modelos'), {
            'acao': 'importar', 'arquivo': SimpleUploadedFile('TCB-OTB-99.docx', formulario_docx()),
        })
        self.modelo = Modelo.objects.get(codigo='TCB-OTB-99')
        self.assertRedirects(r, reverse('checklists:modelo', args=[self.modelo.pk]))
        self.tipo = self.modelo.tipos_ativo.get()
        self.tipo.medidor = 'horimetro'
        self.tipo.save()
        self.ativo = Ativo.objects.create(tipo=self.tipo, nome='Gerador 01', identificacao='GD-4471', marca='Stemac', medidor_atual=120)
        self.client.logout()

    def _post(self, user=None, **extra):
        rev = self.modelo.revisao_vigente
        data = {f'item_{i.pk}': 'ok' for i in rev.itens()}
        data.update({'medidor': '130,5', 'fotos': [_jpeg()], 'assinatura': _assinatura(), 'observacoes': 'Tudo certo.'})
        data.update(extra)
        self.client.force_login(user or self.tec)
        return self.client.post(reverse('checklists:execucao_nova', args=[self.ativo.pk, self.modelo.pk]), data)

    # ---- Importação ----
    def test_importador_le_cabecalho_e_secoes(self):
        caminho = os.path.join(TMP_MEDIA, 'f.docx')
        with open(caminho, 'wb') as f:
            f.write(formulario_docx())
        lido = ler_formulario(caminho)
        self.assertEqual((lido.codigo, lido.revisao), ('TCB-OTB-99', 1))
        self.assertEqual((lido.preparado_por, lido.revisado_por), ('Ana Exemplo', 'Bruno Exemplo'))
        self.assertEqual(lido.data.isoformat(), '2025-05-20')
        self.assertEqual([t for t, _ in lido.secoes], ['Verificação dos EPIs', 'Verificação do equipamento'])
        self.assertEqual(lido.total_itens, 5)

    def test_importacao_cria_modelo_revisao_e_tipo(self):
        rev = self.modelo.revisao_vigente
        self.assertEqual(rev.rotulo, '01')
        self.assertEqual(rev.itens().count(), 5)
        self.assertEqual(self.tipo.nome, 'Gerador diesel')
        # Reimportar a mesma revisão não duplica
        self.client.force_login(self.admin)
        self.client.post(reverse('checklists:modelos'), {'acao': 'importar', 'arquivo': SimpleUploadedFile('a.docx', formulario_docx())})
        self.assertEqual(Revisao.objects.filter(modelo=self.modelo).count(), 1)

    def test_importacao_recusa_arquivo_invalido(self):
        self.client.force_login(self.admin)
        r = self.client.post(reverse('checklists:modelos'), {'acao': 'importar', 'arquivo': SimpleUploadedFile('x.docx', b'nao e word')})
        self.assertContains(r, 'não é um documento Word')

    # ---- Execução ----
    def test_tecnico_executa_e_supervisor_recebe_pdf(self):
        r = self._post()
        ex = Execucao.objects.get()
        self.assertRedirects(r, reverse('checklists:execucao', args=[ex.pk]))
        self.assertEqual(ex.respostas.count(), 5)
        self.assertEqual(ex.executor, self.tec)
        self.assertEqual(ex.medidor_label, '130,5 h')
        self.ativo.refresh_from_db()
        self.assertEqual(float(self.ativo.medidor_atual), 130.5)
        m = mail.outbox[-1]
        self.assertEqual(m.to, ['sup@x.com'])
        self.assertIn('TCB-OTB-99', m.subject)
        self.assertIn('Tiago Técnico', m.subject)
        nome, conteudo, tipo = m.attachments[0]
        self.assertTrue(nome.startswith('TCB-OTB-99_GD-4471_'))
        self.assertTrue(conteudo.startswith(b'%PDF'))

    def test_nok_exige_observacao_e_campos_obrigatorios(self):
        item = self.modelo.revisao_vigente.itens().first()
        r = self._post(**{f'item_{item.pk}': 'nok'})
        self.assertContains(r, 'Descreva a observação')
        r = self._post(fotos=[])
        self.assertContains(r, 'pelo menos uma foto')
        r = self._post(assinatura='')
        self.assertContains(r, 'Falta a assinatura')
        r = self._post(medidor='100')
        self.assertContains(r, 'não pode ser menor que a última registrada')
        self.assertFalse(Execucao.objects.exists())
        r = self._post(**{f'item_{item.pk}': 'nok', f'obs_{item.pk}': 'Óculos riscados'})
        self.assertTrue(Execucao.objects.get().tem_problema)

    def test_item_com_foto_obrigatoria(self):
        item = self.modelo.revisao_vigente.itens().last()
        item.exige_foto = True
        item.save()
        r = self._post()
        self.assertContains(r, 'Falta a foto do item')
        self._post(**{f'foto_{item.pk}': _jpeg('painel.jpg')})
        self.assertEqual(Execucao.objects.get().fotos.filter(item=item).count(), 1)

    def test_tecnico_so_ve_o_proprio(self):
        self._post(user=self.outro)
        dele = Execucao.objects.get()
        self.client.force_login(self.tec)
        self.assertNotContains(self.client.get(reverse('checklists:historico')), 'GD-4471')
        self.assertEqual(self.client.get(reverse('checklists:execucao', args=[dele.pk])).status_code, 404)
        self.assertEqual(self.client.get(reverse('checklists:execucao_pdf', args=[dele.pk])).status_code, 404)
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse('checklists:execucao', args=[dele.pk])).status_code, 200)

    def test_logistica_ve_execucoes_de_todos(self):
        self._post(user=self.outro)
        ex = Execucao.objects.get()
        log = User.objects.create_user(email='log@x.com', password='x', full_name='Lara', is_approved=True, role=User.ROLE_LOGISTICA)
        self.client.force_login(log)
        r = self.client.get(reverse('checklists:historico'))
        self.assertContains(r, 'GD-4471')
        self.assertContains(r, 'Outro Técnico')
        self.assertEqual(self.client.get(reverse('checklists:execucao_pdf', args=[ex.pk])).status_code, 200)
        self.assertEqual(self.client.get(reverse('checklists:modelos')).status_code, 403)

    # ---- Controle de revisão ----
    def test_revisao_usada_fica_bloqueada_e_nova_revisao_preserva_historico(self):
        self._post()
        ex = Execucao.objects.get()
        self.client.force_login(self.admin)
        item = self.modelo.revisao_vigente.itens().first()
        self.client.post(reverse('checklists:modelo', args=[self.modelo.pk]), {
            'acao': 'item', 'id': item.pk, 'texto': 'Mudado', 'tipo_resposta': 'ok_nok', 'observacao': 'se_negativo'})
        item.refresh_from_db()
        self.assertNotEqual(item.texto, 'Mudado')
        self.client.post(reverse('checklists:modelo', args=[self.modelo.pk]), {'acao': 'nova_revisao', 'notas': 'Inclui item de bateria'})
        nova = self.modelo.revisao_vigente
        self.assertEqual(nova.rotulo, '02')
        self.assertEqual(nova.itens().count(), 5)
        ex.refresh_from_db()
        self.assertEqual(ex.revisao.rotulo, '01')
        self.assertTrue(render_execucao_pdf(ex).startswith(b'%PDF'))
        # A nova revisão aceita edição
        nova_item = nova.itens().first()
        self.client.post(reverse('checklists:modelo', args=[self.modelo.pk]), {
            'acao': 'item', 'id': nova_item.pk, 'texto': 'Bateria carregada?', 'tipo_resposta': 'conformidade', 'observacao': 'se_negativo'})
        nova_item.refresh_from_db()
        self.assertEqual((nova_item.texto, nova_item.tipo_resposta), ('Bateria carregada?', 'conformidade'))

    def test_tipos_de_resposta_variados(self):
        rev = self.modelo.revisao_vigente
        secao = rev.secoes.first()
        pct = Item.objects.create(secao=secao, texto='Nível de combustível', tipo_resposta='porcentagem', observacao='nunca', ordem=9)
        num = Item.objects.create(secao=secao, texto='Tensão (V)', tipo_resposta='numero', observacao='nunca', ordem=10)
        conf = Item.objects.create(secao=secao, texto='Aterramento', tipo_resposta='conformidade', ordem=11)
        r = self._post(**{f'item_{pct.pk}': '75', f'item_{num.pk}': '220,5', f'item_{conf.pk}': 'parcial'})
        self.assertContains(r, 'Descreva a observação do item &quot;Aterramento&quot;')
        self._post(**{f'item_{pct.pk}': '75', f'item_{num.pk}': '220,5', f'item_{conf.pk}': 'parcial', f'obs_{conf.pk}': 'Cabo solto'})
        ex = Execucao.objects.get()
        self.assertEqual(ex.respostas.get(item=pct).rotulo, '75%')
        self.assertEqual(ex.respostas.get(item=num).valor, '220.5')
        self.assertTrue(ex.tem_problema)
        self.assertTrue(render_execucao_pdf(ex).startswith(b'%PDF'))

    # ---- Permissões ----
    def test_permissoes(self):
        prof = User.objects.create_user(email='p@x.com', password='x', full_name='P', is_approved=True)
        self.client.force_login(prof)
        self.assertEqual(self.client.get(reverse('checklists:executar')).status_code, 403)
        self.assertEqual(self.client.get(reverse('veiculos:inspecao_nova')).status_code, 403)
        self.client.force_login(self.tec)
        self.assertContains(self.client.get(reverse('checklists:executar')), 'Gerador 01')
        for name in ['tipos', 'ativos', 'modelos']:
            self.assertEqual(self.client.get(reverse(f'checklists:{name}')).status_code, 403, name)
        home = self.client.get(reverse('dashboard:home'))
        self.assertContains(home, 'Pré-uso de equipamento')
        self.assertContains(home, 'Vistoria de veículo')
        self.assertNotContains(home, 'Configurações')
        self.client.force_login(self.admin)
        for name in ['tipos', 'ativos', 'modelos']:
            self.assertContains(self.client.get(reverse(f'checklists:{name}')), 'Formulários')
        self.assertContains(self.client.get(reverse('checklists:modelo', args=[self.modelo.pk])), 'Óculos de segurança?')
